# -*- coding: utf-8 -*-
"""OEM诊断调查表(XLSX)解析器测试（openpyxl构造合成调查表，无需真实OEM文件）"""

import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(
    os.path.dirname(__file__), "..", "..")))

import openpyxl
from openpyxl.styles import Font

from src.business.survey_xlsx_parser import (
    parse_survey_xlsx, save_survey_json,
    _hex_token, _parse_conversion, _map_data_type, _build_request,
    _parse_int, _parse_float)


def _make_service_sheet(wb):
    """服务矩阵sheet: 含支持/不支持/抑制位/删除线/预编程NRC声明各类场景"""
    ws = wb.create_sheet("1_1 Services")
    # 表头（第2行，验证表头不固定在首行也能定位）
    header = ["Service ID", "Services Name", "Supported Services",
              "Sub-functions", "Sub-functions Supported", "SPRMB2"]
    for j, text in enumerate(header, start=1):
        ws.cell(row=2, column=j, value=text)
    rows = [
        # (sid, name, supported, sub, subsup, sprmb, nrc声明)
        ("0x10", "DiagnosticSessionControl", "Y", "0x01 DefaultSession", "Y", "N", ""),
        ("0x10", "", "", "0x02 ProgrammingSession", "Y", "N", "22,13"),  # 预编程NRC
        ("0x10", "", "", "0x03 ExtendedSession", "N", "N", ""),        # 不支持
        ("0x27", "SecurityAccess", "Y", "0x01", "Y", "Y", ""),          # 抑制位变体
        ("0x2E", "WriteDataByIdentifier", "Y", "0x01", "Y", "N", ""),   # 需运行时参数
    ]
    for i, (sid, name, sup, sub, subsup, sprm, nrc) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=sid)
        ws.cell(row=i, column=2, value=name)
        ws.cell(row=i, column=3, value=sup)
        ws.cell(row=i, column=4, value=sub)
        ws.cell(row=i, column=5, value=subsup)
        ws.cell(row=i, column=6, value=sprm)
        if nrc:
            ws.cell(row=i, column=8, value=nrc)
    # 删除线行: TesterPresent整条废弃
    ws.cell(row=8, column=1, value="0x3E").font = Font(strike=True)
    ws.cell(row=8, column=2, value="TesterPresent")
    ws.cell(row=8, column=4, value="0x00")
    ws.cell(row=8, column=5, value="Y")
    return ws


def _make_did_sheet(wb):
    """读写DID列表sheet: 双DID+权限矩阵+换算公式+删除线+多行位描述"""
    ws = wb.create_sheet("4_1 DID")
    header = ["DID Num", "DID Description", "Size", "Data Content",
              "Range, Min", "Range, Max", "Unit", "Conversion",
              "Data Type", "Security Level", "Comments",
              "Application Software(Diagnostic Session)",
              "Boot Sofeware(Diagnostic Session)"]
    for j, text in enumerate(header, start=1):
        ws.cell(row=1, column=j, value=text)
    # 会话权限子表头（表头行下一行，列号须位于Application Software表头之后）
    ws.cell(row=2, column=13, value="0x01")
    ws.cell(row=2, column=14, value="0x03")
    # DID1: VIN（ASCII，多行位描述，App/Boot会话R/RW）
    ws.cell(row=3, column=1, value="0xF190")
    ws.cell(row=3, column=2, value="Vehicle Identification Number")
    ws.cell(row=3, column=3, value="20")
    ws.cell(row=3, column=9, value="ASCII")
    ws.cell(row=3, column=10, value="N")
    ws.cell(row=3, column=13, value="R")
    ws.cell(row=3, column=14, value="RW")
    ws.cell(row=4, column=4, value="Byte0: VIN[0]")
    ws.cell(row=5, column=4, value="Byte1: VIN[1]")
    # DID2: 发动机转速（换算公式，只读，安全等级Level1）
    ws.cell(row=6, column=1, value="0x0100")
    ws.cell(row=6, column=2, value="Engine Speed")
    ws.cell(row=6, column=3, value="2")
    ws.cell(row=6, column=5, value="0")
    ws.cell(row=6, column=6, value="8000")
    ws.cell(row=6, column=7, value="rpm")
    ws.cell(row=6, column=8, value="phy=XX*0.05625")
    ws.cell(row=6, column=9, value="Hex(Unsigned)")
    ws.cell(row=6, column=10, value="Level1")
    ws.cell(row=6, column=13, value="R")
    ws.cell(row=6, column=14, value="R")
    # 删除线DID: 已废弃
    ws.cell(row=7, column=1, value="0x0200").font = Font(strike=True)
    ws.cell(row=7, column=2, value="Deprecated DID")
    return ws


def _make_dtc_sheet(wb):
    ws = wb.create_sheet("3_1 DTC")
    header = ["DTC Display", "DTC Byte", "DTC Meaning", "Faults Attribute"]
    for j, text in enumerate(header, start=1):
        ws.cell(row=1, column=j, value=text)
    ws.cell(row=2, column=1, value="P0100")
    ws.cell(row=2, column=2, value="C00100")
    ws.cell(row=2, column=3, value="Circuit Malfunction\n电路故障")
    ws.cell(row=2, column=4, value="PRESENT")
    # 删除线DTC
    ws.cell(row=3, column=2, value="C00200").font = Font(strike=True)
    ws.cell(row=3, column=3, value="Deprecated fault")
    # Bit Field状态位定义区: 解析到此为止
    ws.cell(row=4, column=1, value="Bit Field 0-7 definition")
    ws.cell(row=5, column=2, value="C00300")   # 不应被解析
    return ws


def _make_routine_sheet(wb):
    ws = wb.create_sheet("4_3 Routine")
    header = ["RoutineDID", "DID Description", "RoutineControlType",
              "Subfuction", "Security Leve"]
    for j, text in enumerate(header, start=1):
        ws.cell(row=1, column=j, value=text)
    ws.cell(row=2, column=1, value="0x0203")
    ws.cell(row=2, column=2, value="PreProgramming Check")
    ws.cell(row=2, column=3, value="0x01 start")
    ws.cell(row=2, column=4, value="Y")
    ws.cell(row=2, column=5, value="Level1")
    return ws


def _make_workbook_path(tmp):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _make_service_sheet(wb)
    _make_did_sheet(wb)
    _make_dtc_sheet(wb)
    _make_routine_sheet(wb)
    path = os.path.join(tmp, "survey.xlsx")
    wb.save(path)
    return path


class TestHelpers(unittest.TestCase):
    """纯函数辅助解析逻辑"""

    def test_hex_token(self):
        self.assertEqual(_hex_token("0x01 DefaultSession"), ("01", "DefaultSession"))
        self.assertEqual(_hex_token("01 DefaultSession"), ("01", "DefaultSession"))
        self.assertEqual(_hex_token("0x01Name"), ("01", "Name"))
        self.assertEqual(_hex_token("多行\n0x0A"), (None, "多行"))  # hex不在首行→无匹配
        # 3位hex后跟数字属非法写法（实测行为: 匹配首字符"0"，剩余作文本）
        self.assertEqual(_hex_token("0x0123"), ("0", "x0123"))

    def test_parse_conversion(self):
        self.assertEqual(_parse_conversion("phy=XX*0.05625"), (0.05625, 0.0))
        self.assertEqual(_parse_conversion("phy=XX*0.01+3"), (0.01, 3.0))
        self.assertEqual(_parse_conversion("XX"), (1.0, 0.0))
        self.assertEqual(_parse_conversion(""), (1.0, 0.0))

    def test_map_data_type(self):
        self.assertEqual(_map_data_type("Hex(Unsigned)"), "uint")
        self.assertEqual(_map_data_type("Signed"), "int")
        self.assertEqual(_map_data_type("ASCII"), "ascii")
        self.assertEqual(_map_data_type("BCD"), "raw")
        self.assertEqual(_map_data_type(""), "raw")

    def test_build_request(self):
        self.assertEqual(_build_request(0x10, "03"), "10 03")
        self.assertEqual(_build_request(0x85, "02"), "85 02")
        self.assertEqual(_build_request(0x28, "00"), "28 00 03")
        self.assertEqual(_build_request(0x19, "02"), "19 02 FF")
        self.assertEqual(_build_request(0x19, "0A"), "19 0A")
        self.assertEqual(_build_request(0x14, "FF"), "14 FF FF FF")
        self.assertIsNone(_build_request(0x2E, "01"))  # 需运行时参数
        self.assertIsNone(_build_request(0x22, ""))

    def test_parse_int_float(self):
        self.assertEqual(_parse_int("20 bytes"), 20)
        self.assertIsNone(_parse_int("abc"))
        self.assertEqual(_parse_float("0.5"), 0.5)
        self.assertIsNone(_parse_float("-"))
        self.assertIsNone(_parse_float(""))


class TestParseSurveyXlsx(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = _make_workbook_path(cls.tmp.name)
        cls.data = parse_survey_xlsx(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_services_generated(self):
        """服务条目: 支持的子功能各生成一条完整请求"""
        requests = [s["request"] for s in self.data["services"]]
        self.assertIn("10 01", requests)
        self.assertIn("10 02", requests)
        self.assertIn("27 01", requests)

    def test_unsupported_subfunction_excluded(self):
        """子功能不支持(N)不生成用例"""
        requests = [s["request"] for s in self.data["services"]]
        self.assertNotIn("10 03", requests)

    def test_suppress_positive_response_variant(self):
        """SPRMB2=Y → 额外生成+0x80抑制正响应变体"""
        names = [s["name"] for s in self.data["services"]]
        requests = [s["request"] for s in self.data["services"]]
        self.assertIn("27 81", requests)
        self.assertIn("Suppress", "".join(names))

    def test_runtime_parameter_service_skipped_with_reason(self):
        """需运行时参数的服务记录跳过原因"""
        skipped = {name: reason for name, reason in self.data["_skipped"]}
        self.assertIn("0x2E WriteDataByIdentifier", skipped)
        self.assertIn("运行时参数", skipped["0x2E WriteDataByIdentifier"])

    def test_struck_service_row_excluded(self):
        """删除线服务（TesterPresent）整条剔除"""
        requests = [s["request"] for s in self.data["services"]]
        self.assertNotIn("3E 00", requests)

    def test_preprogramming_flag(self):
        """10 02行NRC声明含0x22 → 输出预编程条件检查配置"""
        pre = self.data.get("preprogramming")
        self.assertIsNotNone(pre)
        self.assertTrue(pre["required"])
        self.assertEqual(pre["check_request"], "31 01 02 03")

    def test_did_parsed_with_access_and_scaling(self):
        """DID条目: 读写权限归纳 + 换算公式/类型映射"""
        dids = {d["id"]: d for d in self.data["dids"]}
        self.assertIn("0xF190", dids)
        self.assertIn("0x0100", dids)

        vin = dids["0xF190"]
        self.assertEqual(vin["data_type"], "ascii")
        self.assertEqual(vin["access"], "rw")   # R + RW → rw
        self.assertEqual(vin["data_length"], 20)
        # 多行位描述合并
        self.assertIn("VIN[0]", vin["description"])
        self.assertIn("VIN[1]", vin["description"])

        speed = dids["0x0100"]
        self.assertEqual(speed["data_type"], "uint")
        self.assertEqual(speed["access"], "r")  # 仅R
        self.assertAlmostEqual(speed["scaling"], 0.05625)
        self.assertEqual(speed["offset"], 0.0)
        self.assertEqual(speed["unit"], "rpm")
        self.assertEqual(speed["min_val"], 0.0)
        self.assertEqual(speed["max_val"], 8000.0)
        self.assertEqual(speed["security"], "Level1")

    def test_struck_did_excluded(self):
        ids = [d["id"] for d in self.data["dids"]]
        self.assertNotIn("0x0200", ids)

    def test_dtc_parsed(self):
        """DTC条目: 双语描述合并为单行"""
        dtcs = self.data["dtcs"]
        self.assertEqual(len(dtcs), 1)
        self.assertEqual(dtcs[0]["dtc_id"], "0xC00100")
        self.assertIn("Circuit Malfunction", dtcs[0]["description"])
        self.assertIn("PRESENT", dtcs[0]["description"])

    def test_dtc_parsing_stops_at_bit_field(self):
        """Bit Field区后的DTC不解析"""
        dtc_ids = [d["dtc_id"] for d in self.data["dtcs"]]
        self.assertNotIn("0xC00300", dtc_ids)

    def test_struck_dtc_excluded(self):
        dtc_ids = [d["dtc_id"] for d in self.data["dtcs"]]
        self.assertNotIn("0xC00200", dtc_ids)

    def test_routine_parsed(self):
        """例程条目: 31 <控制类型> <例程ID>"""
        requests = [s["request"] for s in self.data["services"]]
        self.assertIn("31 01 0203", requests)
        routine = next(s for s in self.data["services"]
                        if s["request"] == "31 01 0203")
        self.assertIn("PreProgramming Check", routine["description"])
        self.assertIn("Level1", routine["description"])

    def test_services_deduplicated_by_request(self):
        reqs = [s["request"] for s in self.data["services"]]
        self.assertEqual(len(reqs), len(set(reqs)))

    def test_save_survey_json_strips_internal_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.json")
            save_survey_json(self.data, path)
            with open(path, "r", encoding="utf-8") as f:
                saved = json.load(f)
        for key in saved:
            self.assertFalse(key.startswith("_"),
                             f"内部字段 {key} 不应写入输出JSON")
        self.assertIn("services", saved)
        self.assertIn("dids", saved)
        self.assertIn("dtcs", saved)

    def test_parse_nonexistent_file_raises(self):
        with self.assertRaises(Exception):
            parse_survey_xlsx(os.path.join("Z:\\no", "such", "file.xlsx"))


if __name__ == "__main__":
    unittest.main()
